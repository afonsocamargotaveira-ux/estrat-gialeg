"""Preenche o modelo 'Estratégia para pesquisa de proposições no Sileg'
mantendo 100% da formatação do arquivo original (.docx)."""
import copy
import io
import re
from pathlib import Path

from docx import Document
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from PIL import Image

MODELO = Path(__file__).with_name("modelo_estrategia.docx")
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

LARGURA_MAX_CM = 14.0   # largura útil (17 cm) menos o recuo de 3 cm da lista
ALTURA_MAX_CM = 20.0


# ---------- utilidades ----------
def _set_text(par, texto):
    """Põe o texto no 1º run (preservando sua formatação) e apaga os demais."""
    runs = par.runs
    runs[0].text = texto
    for r in runs[1:]:
        r._element.getparent().remove(r._element)
    for pe in par._element.findall(W + "proofErr"):
        par._element.remove(pe)


def _limpa_quebras_render(el):
    for x in el.iter(W + "lastRenderedPageBreak"):
        x.getparent().remove(x)


def _linhas(texto):
    return [l.strip() for l in (texto or "").splitlines() if l.strip()]


def _juntar_pt(itens):
    itens = list(itens)
    if len(itens) <= 1:
        return "".join(itens)
    return ", ".join(itens[:-1]) + " e " + itens[-1]


def _substitui_em_wt(elemento, novo):
    for t in elemento.iter(W + "t"):
        if t.text and "Mulher" in t.text:
            t.text = t.text.replace("Mulher", novo)


# ---------- geração ----------
def gerar(tema, data_atualizacao, protocolo, sequencias,
          lista_exclusao, nao_excluir, temas_excluir, imagens_estrutura=()):
    """
    sequencias: lista de dicts {campos: [str], texto: str}
    imagens_estrutura: lista de bytes (imagens de como ficará a estrutura da pesquisa)
    lista_exclusao: str (proposições separadas por vírgula)
    nao_excluir / temas_excluir: str, um termo por linha
    Retorna os bytes do .docx.
    """
    doc = Document(str(MODELO))
    body = doc.element.body
    pars = {i: Paragraph(e, doc._body) for i, e in enumerate(body) if e.tag == W + "p"}
    el = list(body)                                      # índice igual ao da análise

    # 0/1 – data da atualização
    _set_text(pars[1], data_atualizacao)

    # 2 – título (tema)
    pars[2].runs[0].text = tema

    # cabeçalho e rodapé
    sec = doc.sections[0]
    _substitui_em_wt(sec.header._element, tema)
    _substitui_em_wt(sec.footer._element, tema)

    # 13 – número de sequências
    for r in pars[13].runs:
        if "______" in r.text:
            r.text = str(len(sequencias))

    # 14..19 – sequências de termos (cabeçalho, texto e prints)
    mod_titulo = copy.deepcopy(el[14])
    mod_corpo = copy.deepcopy(el[17])
    _limpa_quebras_render(mod_corpo)
    for r in list(mod_corpo.findall(W + "r")) + list(mod_corpo.findall(W + "proofErr")):
        mod_corpo.remove(r)
    ancora = el[20]

    def novo_corpo():
        return copy.deepcopy(mod_corpo)

    for i, seq in enumerate(sequencias, 1):
        t = copy.deepcopy(mod_titulo)
        pt = Paragraph(t, doc._body)
        campos = _juntar_pt(seq["campos"]) or "ementa, explicação da ementa e indexação"
        pt.runs[0].text = str(i)
        pt.runs[1].text = f"ª sequência (marcar {campos}):"
        ancora.addprevious(t)

        for linha in _linhas(seq["texto"]):
            c = novo_corpo()
            Paragraph(c, doc._body).add_run(linha)
            ancora.addprevious(c)

        sep = novo_corpo()                                # linha em branco entre sequências
        ancora.addprevious(sep)

    for img in imagens_estrutura:                         # imagem(ns) da estrutura, após as sequências
        c = novo_corpo()
        pc = Paragraph(c, doc._body)
        with Image.open(io.BytesIO(img)) as im:
            w_px, h_px = im.size
        w_cm = min(LARGURA_MAX_CM, w_px / 96 * 2.54)
        h_cm = w_cm * h_px / w_px
        if h_cm > ALTURA_MAX_CM:
            w_cm *= ALTURA_MAX_CM / h_cm
        pc.add_run().add_picture(io.BytesIO(img), width=Cm(w_cm))
        ancora.addprevious(c)

    for i in (14, 15, 16, 17, 18, 19):                      # marcadores do modelo
        body.remove(el[i])

    # 39 – lista de exclusão (horizontal, separada por vírgulas)
    texto_excl = ", ".join(_linhas(lista_exclusao.replace(",", "\n")))
    _set_text(pars[39], texto_excl)

    # 48 – protocolo Prisma
    for r in pars[48].runs:
        if "____" in r.text:
            r.text = protocolo

    # 76 e 79 – listas de termos (um parágrafo por linha)
    for idx, valor in ((76, nao_excluir), (79, temas_excluir)):
        linhas = _linhas(valor)
        modelo = el[idx]
        if not linhas:
            _set_text(pars[idx], "")
            continue
        for linha in linhas:
            c = copy.deepcopy(modelo)
            _set_text(Paragraph(c, doc._body), linha)
            modelo.addprevious(c)
        body.remove(modelo)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def nome_arquivo(tema):
    s = re.sub(r"[^\w\- ]", "", tema, flags=re.UNICODE).strip().replace(" ", "_")
    return f"Estrategia_{s or 'tema'}_Sileg.docx"
